from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = [
    "Riga Excel",
    "Campo",
    "Valore ricevuto",
    "Codice Errore",
    "Livello",
    "Descrizione",
    "Valori ammessi / Regola",
]

AIFA_HEADERS = [
    "Riga Excel",
    "Campo",
    "Valore originale",
    "Valore AIFA",
    "AIC",
    "Nome AIFA",
    "Azienda AIFA",
    "ATC AIFA",
    "Fonte",
    "Metodo",
    "Confidence",
]

AIFA_CHECK_HEADERS = [
    "Riga Excel",
    "AIC",
    "Match AIFA",
    "Nome ricevuto",
    "Nome AIFA",
    "Principio ricevuto",
    "Principio AIFA",
    "ATC7 ricevuto",
    "ATC AIFA",
    "Fornitore ricevuto",
    "Azienda titolare AIFA",
    "Stato amministrativo",
    "Azione",
]

SUPPLIER_HEADERS = [
    "Riga Excel",
    "AIC",
    "Nome AIFA",
    "Fornitore ricevuto",
    "Azienda titolare AIFA",
    "Azione",
    "Nota",
]

AIFA_UPC_HEADERS = [
    "Riga Excel",
    "AIC",
    "Nome Commerciale",
    "UPC ricevuto",
    "Unità Posologiche AIFA",
    "Fonte unità",
    "Rapporto UPC/AIFA",
    "Esito",
    "Nota",
]


OPERATOR_CHANGE_HEADERS = [
    "Data/Ora",
    "Riga Excel",
    "Campo",
    "Valore precedente",
    "Valore nuovo",
    "Azione",
]

TRANSFORMATION_HEADERS = [
    "Riga Excel",
    "Campo",
    "Valore originale",
    "Valore normalizzato",
    "Motivo",
]


def _fit_columns(ws, minimum=12, maximum=55):
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(text))
        ws.column_dimensions[letter].width = min(max(max_length + 2, minimum), maximum)


def _format_header(row, fill):
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _sheet_from_records(wb, name, headers, records, title_fill):
    ws = wb.create_sheet(name)
    ws.append(headers)
    _format_header(ws[1], title_fill)
    for item in records:
        ws.append([item.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _fit_columns(ws)
    return ws


def build_validation_report(source_filename, result):
    output = BytesIO()
    wb = Workbook()

    ws = wb.active
    ws.title = "Riepilogo"

    title_fill = PatternFill("solid", fgColor="1F4E78")
    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    ko_fill = PatternFill("solid", fgColor="FCE4D6")

    ws["A1"] = "Report Validazione - Listino Farmaci"
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws.merge_cells("A1:B1")

    ds = result.get("summary", {})
    summary = [
        ("File analizzato", source_filename),
        ("Data controllo", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Fornitore", ds.get("supplier", "-")),
        ("Numero fornitori rilevati", ds.get("supplier_count", "-")),
        ("Riga intestazioni rilevata", result.get("header_row", "-")),
        ("Righe dati analizzate", result["rows_count"]),
        ("Articoli con AIC", ds.get("aic_count", "-")),
        ("AIC riconosciuti da AIFA", ds.get("aifa_exact_matches", 0)),
        ("AIC non trovati da AIFA", ds.get("aifa_missing_matches", 0)),
        ("Campi completati da AIFA", ds.get("aifa_enrichment_count", 0)),
        ("Proposte Fornitore da Azienda titolare", ds.get("supplier_proposal_count", 0)),
        ("Principi attivi AIFA non disponibili", ds.get("without_active_count", "-")),
        ("ATC7 ancora mancanti", ds.get("without_atc_count", "-")),
        ("UPC coerenti con AIFA", ds.get("aifa_upc_ok", 0)),
        ("UPC da verificare vs AIFA", ds.get("aifa_upc_warning", 0)),
        ("UPC non confrontabili AIFA", ds.get("aifa_upc_unchecked", 0)),
        ("Normalizzazioni applicate", ds.get("normalization_count", 0)),
        ("Modifiche operatore", len(result.get("operator_changes", []))),
        ("Errori bloccanti", result["blocking_count"]),
        ("Warning", result["warning_count"]),
        ("Informazioni", result["info_count"]),
        ("Esito", "VALIDATO" if result["is_valid"] else "NON VALIDATO"),
    ]

    for row_idx, (label, value) in enumerate(summary, start=3):
        ws.cell(row_idx, 1, label).font = Font(bold=True)
        ws.cell(row_idx, 2, value)

    outcome_row = 3 + len(summary) - 1
    ws.cell(outcome_row, 2).fill = ok_fill if result["is_valid"] else ko_fill
    _fit_columns(ws)

    err_ws = _sheet_from_records(wb, "Anomalie", HEADERS, result["issues"], title_fill)
    for row in err_ws.iter_rows(min_row=2):
        level = row[4].value
        if level == "BLOCCANTE":
            row[4].fill = PatternFill("solid", fgColor="F4CCCC")
        elif level == "WARNING":
            row[4].fill = PatternFill("solid", fgColor="FFF2CC")
        elif level == "INFO":
            row[4].fill = PatternFill("solid", fgColor="D9EAF7")

    _sheet_from_records(wb, "Controllo anagrafico AIFA", AIFA_CHECK_HEADERS, result.get("aifa_checks", []), title_fill)
    _sheet_from_records(wb, "Arricchimenti AIFA", AIFA_HEADERS, result.get("aifa_enrichments", []), title_fill)
    _sheet_from_records(wb, "Proposte Fornitore AIFA", SUPPLIER_HEADERS, result.get("supplier_proposals", []), title_fill)

    upc_ws = _sheet_from_records(wb, "Controllo UPC AIFA", AIFA_UPC_HEADERS, result.get("aifa_upc_checks", []), title_fill)
    # Esito ora è la colonna H (indice 8 Excel)
    for row in upc_ws.iter_rows(min_row=2):
        esito = row[7].value
        if esito == "COERENTE":
            row[7].fill = PatternFill("solid", fgColor="E2F0D9")
        elif esito == "DA VERIFICARE":
            row[7].fill = PatternFill("solid", fgColor="FFF2CC")

    _sheet_from_records(wb, "Normalizzazioni", TRANSFORMATION_HEADERS, result.get("transformations", []), title_fill)
    _sheet_from_records(wb, "Modifiche operatore", OPERATOR_CHANGE_HEADERS, result.get("operator_changes", []), title_fill)

    records = result.get("normalized_records", [])
    if records:
        data_ws = wb.create_sheet("Dataset normalizzato")
        headers = list(records[0].keys())
        data_ws.append(headers)
        _format_header(data_ws[1], title_fill)
        for record in records:
            data_ws.append([record.get(header) for header in headers])
        data_ws.freeze_panes = "A2"
        data_ws.auto_filter.ref = data_ws.dimensions
        _fit_columns(data_ws, maximum=45)

    wb.save(output)
    output.seek(0)
    return output.getvalue()


def build_normalized_workbook(workbook_data, schema):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = schema.get("sheet_name", "Tracciato")

    headers = [column["name"] for column in schema["columns"]]
    ws.append(headers)

    title_fill = PatternFill("solid", fgColor="1F4E78")
    _format_header(ws[1], title_fill)

    for record in workbook_data.records:
        ws.append([record.get(header) for header in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _fit_columns(ws, maximum=45)

    wb.save(output)
    output.seek(0)
    return output.getvalue()
