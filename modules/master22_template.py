from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


GUIDANCE = {
    "Fornitore": "Indicare la ragione sociale del fornitore",
    "AIC": "Inserire il codice AIC unico dell'articolo",
    "Codice Fornitore": "In assenza di codice AIC indicare il codice interno ad uso del fornitore per identificare l'articolo",
    "Nome Commerciale": "Nome Commerciale dell'articolo",
    "Principio Attivo": "Indicare il principio attivo",
    "Forma Farmaceutica": "Indicare la forma farmaceutica (es: Pezzo, Flacone…)",
    "Materiale Pericoloso": "Flag Y / N per indicare se si tratta di materiale pericoloso",
    "Stupefacente": "Indicare il campo secondo la lista proposta",
    "ATC7": "Riportare la descrizione del campo ATC7",
    "ATC9": "Riportare la descrizione del campo ATC9",
    "Fala / Lasa": "Flag Y / N per indicare se si tratta di un articolo Fala / Lasa",
    "Gruppo di Stivaggio": "Campo ad imputazione del Buyer Gksd ProCure secondo l'elenco prestabilito per indicare il gruppo di stivaggio dell'articolo.",
    "Temperatura di Stivaggio": "Indicare la fascia di temperatura per lo stivaggio in magazzino.",
    "Prezzo Unitario": "Prezzo unitario (es: costo singola pasticca o unità indivisibile di cui si compone l'articolo)",
    "Prezzo Confezione": "Prezzo a Confezione (moltiplicatore dell'unità per confezione per le singole unità indivisibili presenti nel blister)",
    "UPC": "Unità per confezione",
    "Minimo Movimentabile": "Numero minimo di unità indivisibili movimentabili. Gli ordini successivi devono rispettarne i multipli.",
    "IVA": "Indicare l'aliquota IVA alla quale è soggetto l'articolo in essere",
    "Note": "Campo note per indicazioni generiche sulle condizioni di vendita e/o acquisto dell'articolo",
    "X": "Indicare la lunghezza della confezione, espressa in centimetri (cm), considerando l'ingombro massimo esterno.",
    "Y": "Indicare la larghezza della confezione, espressa in centimetri (cm), considerando l'ingombro massimo esterno.",
    "Z": "Indicare l'altezza della confezione, espressa in centimetri (cm), considerando l'ingombro massimo esterno.",
    "Data Validità Farmaco": "Indicare la data fino alla quale l'AIC può rimanere in commercio. Formato consigliato: GG/MM/AAAA.",
}


def ensure_master_template(schema: dict, root: str | Path = ".") -> Path:
    reference = schema.get("template_reference", {})
    relative = reference.get("bundled_name")
    if not relative:
        raise ValueError("Template master non configurato nello schema.")

    path = Path(root) / relative
    if path.exists():
        return path

    path.parent.mkdir(parents=True, exist_ok=True)
    _build_template(path, schema)
    return path


def _build_template(path: Path, schema: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = schema.get("sheet_name", "Tracciato")

    headers = [column["name"] for column in schema["columns"]]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    guide_fill = PatternFill("solid", fgColor="D9EAF7")

    for col_idx, header in enumerate(headers, start=1):
        guide_cell = ws.cell(1, col_idx, GUIDANCE.get(header, ""))
        guide_cell.fill = guide_fill
        guide_cell.alignment = Alignment(wrap_text=True, vertical="top")

        header_cell = ws.cell(2, col_idx, header)
        header_cell.fill = header_fill
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        width = 18
        if header in {"Nome Commerciale", "Principio Attivo", "Forma Farmaceutica", "Note"}:
            width = 24
        if header in {"Gruppo di Stivaggio", "Temperatura di Stivaggio"}:
            width = 26
        if header == "Data Validità Farmaco":
            width = 22
        ws.column_dimensions[header_cell.column_letter].width = width

    ws.row_dimensions[1].height = 75
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{ws.cell(2, len(headers)).coordinate}"

    lists_ws = wb.create_sheet("Elenchi")
    list_defs = [
        ("IVA", [f"{float(v) * 100:g}%" for v in schema.get("lists", {}).get("IVA", [])]),
        ("Y / N", schema.get("lists", {}).get("YN", [])),
        ("Gruppo di Stoccaggio", schema.get("lists", {}).get("GRUPPO_STIVAGGIO", [])),
        ("Temperatura", schema.get("lists", {}).get("TEMPERATURA_STIVAGGIO", [])),
        ("Stupefacente", schema.get("lists", {}).get("STUPEFACENTE", [])),
    ]
    for col_idx, (title, values) in enumerate(list_defs, start=1):
        cell = lists_ws.cell(1, col_idx, title)
        cell.font = Font(bold=True)
        for row_idx, value in enumerate(values, start=2):
            lists_ws.cell(row_idx, col_idx, value)
        lists_ws.column_dimensions[cell.column_letter].width = 34

    header_to_col = {name: idx for idx, name in enumerate(headers, start=1)}
    validation_map = {
        "Materiale Pericoloso": (2, len(schema.get("lists", {}).get("YN", [])), "B"),
        "Fala / Lasa": (2, len(schema.get("lists", {}).get("YN", [])), "B"),
        "Gruppo di Stivaggio": (2, len(schema.get("lists", {}).get("GRUPPO_STIVAGGIO", [])), "C"),
        "Temperatura di Stivaggio": (2, len(schema.get("lists", {}).get("TEMPERATURA_STIVAGGIO", [])), "D"),
        "Stupefacente": (2, len(schema.get("lists", {}).get("STUPEFACENTE", [])), "E"),
        "IVA": (2, len(schema.get("lists", {}).get("IVA", [])), "A"),
    }
    for field, (start, count, source_col) in validation_map.items():
        if field not in header_to_col or count <= 0:
            continue
        end = start + count - 1
        dv = DataValidation(type="list", formula1=f"='Elenchi'!${source_col}${start}:${source_col}${end}", allow_blank=False)
        ws.add_data_validation(dv)
        target_col = ws.cell(2, header_to_col[field]).column_letter
        dv.add(f"{target_col}3:{target_col}10000")

    if "Data Validità Farmaco" in header_to_col:
        validity_col = ws.cell(2, header_to_col["Data Validità Farmaco"]).column_letter
        for row_idx in range(3, 10001):
            ws[f"{validity_col}{row_idx}"].number_format = "DD/MM/YYYY"

    wb.save(path)
