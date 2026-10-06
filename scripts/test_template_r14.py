from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from modules.master22_template import GUIDANCE, ensure_master_template
from modules.schema_loader import load_schema


def test_r14_template_generation():
    schema = load_schema()
    with TemporaryDirectory() as tmp:
        path = ensure_master_template(schema, root=tmp)
        assert Path(path).exists()
        wb = load_workbook(path)
        ws = wb[schema.get("sheet_name", "Tracciato")]
        headers = [ws.cell(2, idx).value for idx in range(1, 24)]
        assert len(headers) == 23
        assert headers[4] == "Principio Attivo"
        assert headers[5] == "Forma Farmaceutica"
        assert headers[-4:] == ["X", "Y", "Z", "Data Validità Farmaco"]
        assert ws.cell(1, 5).value == GUIDANCE["Principio Attivo"]
        assert ws.cell(1, 6).value == GUIDANCE["Forma Farmaceutica"]
        assert ws.cell(1, 23).value == GUIDANCE["Data Validità Farmaco"]


if __name__ == "__main__":
    test_r14_template_generation()
    print("R14 template tests: OK")
